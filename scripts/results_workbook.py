"""Offline Excel workbook with the results of one real-data run (no Fabric connection needed).

  uv run python scripts/results_workbook.py --results /tmp/r/results --out excel/Bilanzwerk_Ergebnisse.xlsx

The live monthly report (excel/README.md) needs the Power BI semantic model; this workbook holds
the run's result files as values with native Excel charts, so the results can be opened anywhere.
Every value is copied from the run's files.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HEAD = PatternFill("solid", fgColor="1F5AA6")
WHITE = Font(color="FFFFFF", bold=True)
TITLE = Font(size=14, bold=True)
NOTE = Font(size=9, italic=True, color="52514E")
MODELS = {
    "M0_zscore_rule": "M0 Z''-Regel",
    "M1_logit_z": "M1 Z'' (Logit)",
    "M2_logit_z_wc": "M2 Z'' + Working Capital",
    "M2r_logit_z_wc_relative": "M2r sektorrelativ",
    "S1_discrete_time_hazard": "S1 Hazard",
    "M3_lightgbm_all": "M3 LightGBM, alle KPIs",
}


def put_table(ws, df: pd.DataFrame, row: int, name: str, formats: dict[str, str] | None = None) -> tuple[int, int]:
    """Writes df at `row` as an Excel table; returns (first data row, last data row)."""
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=row, column=j, value=col)
    for i, rec in enumerate(df.itertuples(index=False), start=row + 1):
        for j, val in enumerate(rec, start=1):
            c = ws.cell(row=i, column=j, value=None if pd.isna(val) else val)
            fmt = (formats or {}).get(df.columns[j - 1])
            if fmt:
                c.number_format = fmt
    last = row + len(df)
    ref = f"A{row}:{get_column_letter(len(df.columns))}{last}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for j, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(str(v)) for v in df.iloc[:, j - 1].head(50))) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 10), 60)
    return row + 1, last


def header(ws, title: str, note: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].font = NOTE


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--results", type=Path, required=True)
    p.add_argument("--out", type=Path, default=Path("excel/Bilanzwerk_Ergebnisse.xlsx"))
    args = p.parse_args()
    r = args.results
    run_txt = (r / "RUN.txt").read_text()
    run_url = run_txt.splitlines()[0].split(": ", 1)[1]
    run = run_url.rsplit("/", 1)[-1]
    note = f"Quelle: Lauf auf echten Daten {run} ({run_url}), Branch real-data-results, Ordner results/."
    a = json.loads((r / "analysis/answer.json").read_text())

    wb = Workbook()
    ws = wb.active
    ws.title = "Antwort"
    header(ws, "Bilanzwerk: Antwort auf H1\u2013H4", note)
    rows = []
    for h in ("H1", "H2", "H3"):
        d = a[h]
        rows.append(
            [
                h,
                d["hypothesis"],
                "bestätigt" if d["held"] else "nicht bestätigt",
                d["estimate"],
                d["ci_low"],
                d["ci_high"],
                d["text_de"],
            ]
        )
    h4 = a["H4"]
    rows.append(
        [
            "H4",
            h4["hypothesis"],
            "bestätigt" if h4["held"] else "nicht bestätigt",
            h4["total_eur_bn"],
            None,
            None,
            h4["text_de"],
        ]
    )
    ch = a["champion_check"]
    rows.append(
        [
            "Champion",
            "PR-AUC M3 minus M2",
            "M3 bleibt" if ch["ci_low"] > 0 else "M2",
            ch["estimate"],
            ch["ci_low"],
            ch["ci_high"],
            "ADR-007",
        ]
    )
    df = pd.DataFrame(rows, columns=["Nr", "Hypothese", "Ergebnis", "Schätzer", "95%-KI unten", "95%-KI oben", "Text"])
    put_table(ws, df, 4, "Antwort", {"Schätzer": "0.000", "95%-KI unten": "0.000", "95%-KI oben": "0.000"})
    ws.column_dimensions["G"].width = 110
    for row in ws.iter_rows(min_row=5, max_row=9, min_col=7, max_col=7):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A12"] = "Einheiten: H1 in Monaten (CCC minus Z''), H2/H3 und Champion in PR-AUC-Punkten, H4 in Mrd. EUR."
    ws["A12"].font = NOTE

    # H4 sectors with chart
    ws = wb.create_sheet("H4 Sektoren")
    header(ws, "Working-Capital-Potenzial je Sektor (Mrd. EUR)", note)
    sec = pd.read_csv(r / "analysis/wc_opportunity_by_sector.csv")
    sec = sec.rename(
        columns={
            "sector_de": "Sektor",
            "sector_en": "Sector",
            "cash_opportunity_eur_bn": "Median-Ziel Mrd. EUR",
            "stretch_eur_bn": "Bestes Quartil Mrd. EUR",
            "companies": "Unternehmen",
        }
    )
    first, last = put_table(
        ws, sec, 4, "Sektoren", {"Median-Ziel Mrd. EUR": "#,##0.0", "Bestes Quartil Mrd. EUR": "#,##0.0"}
    )
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Potenzial bis zum Sektormedian (Mrd. EUR)"
    chart.add_data(Reference(ws, min_col=3, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    chart.y_axis.majorGridlines = None
    chart.legend = None
    chart.height, chart.width = 9, 18
    chart.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(chart, "H4")

    ws = wb.create_sheet("H4 Top 20")
    header(ws, "Die 20 Unternehmen mit dem größten Potenzial (letztes Geschäftsjahr)", note)
    top = pd.read_csv(r / "analysis/wc_opportunity_top20.csv")
    top["cash_opportunity_eur"] = top["cash_opportunity_eur"] / 1e9
    top["cash_opportunity_stretch_eur"] = top["cash_opportunity_stretch_eur"] / 1e9
    top = top.rename(
        columns={
            "company_name": "Unternehmen",
            "sector_de": "Sektor",
            "fiscal_year": "GJ",
            "dso": "DSO",
            "dio": "DIO",
            "dpo": "DPO",
            "cash_opportunity_eur": "Median-Ziel Mrd. EUR",
            "cash_opportunity_stretch_eur": "Bestes Quartil Mrd. EUR",
        }
    )
    put_table(
        ws,
        top,
        4,
        "Top20",
        {c: "0.0" for c in ("DSO", "DIO", "DPO")} | {"Median-Ziel Mrd. EUR": "0.00", "Bestes Quartil Mrd. EUR": "0.00"},
    )

    # models
    ws = wb.create_sheet("Modelle")
    header(ws, "Rollierender Backtest 2014\u20132025, gepoolt", note)
    pooled = pd.read_csv(r / "analysis/backtest_pooled.csv")
    pooled = pooled[pooled.model.isin(MODELS)].sort_values("pr_auc", ascending=False)
    pooled.insert(0, "Modell", pooled.model.map(MODELS))
    pooled = pooled.drop(columns="model").rename(
        columns={
            "n": "Beobachtungen",
            "n_pos": "Insolvenzen",
            "pr_auc": "PR-AUC",
            "recall_top5": "Recall Top 5 %",
            "brier": "Brier",
        }
    )
    put_table(
        ws,
        pooled,
        4,
        "Pooled",
        {
            "PR-AUC": "0.000",
            "Recall Top 5 %": "0%",
            "Brier": "0.0000",
            "Beobachtungen": "#,##0",
            "Insolvenzen": "#,##0",
        },
    )

    by = pd.read_csv(r / "analysis/backtest_by_year.csv")
    wide = by[by.model.isin(MODELS)].pivot(index="test_year", columns="model", values="pr_auc")
    wide = (
        wide[[m for m in MODELS if m in wide.columns]]
        .rename(columns=MODELS)
        .reset_index()
        .rename(columns={"test_year": "Testjahr"})
    )
    start = 13
    ws.cell(row=start - 1, column=1, value="PR-AUC je Testjahr").font = Font(bold=True)
    first, last = put_table(ws, wide, start, "ByYear", {c: "0.000" for c in wide.columns if c != "Testjahr"})
    line = LineChart()
    line.title = "PR-AUC je Testjahr"
    line.add_data(
        Reference(ws, min_col=2, max_col=len(wide.columns), min_row=first - 1, max_row=last), titles_from_data=True
    )
    line.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    line.height, line.width = 9, 20
    ws.add_chart(line, f"A{last + 3}")

    # H1 lead times
    ws = wb.create_sheet("H1 Vorlauf")
    header(ws, "H1: Vorlauf der roten Flaggen vor der Insolvenz", note)
    h1 = a["H1"]
    s = h1["sensitivity_valid_obs"]
    rows = [
        ["Insolvenzen mit Kennzahlen", h1["n_companies"], h1["n_companies"]],
        ["Median-Vorlauf (Monate)", h1["median_ccc"], h1["median_z"]],
        ["Median-Vorlauf nur geflaggte (Monate)", h1["median_ccc_flagged_only"], h1["median_z_flagged_only"]],
        ["Anteil bis zum Ereignis rot", 1 - h1["share_never_red_ccc"], 1 - h1["share_never_red_z"]],
        ["Anteil im letzten Quartal berechenbar", h1["share_ccc_available_last"], h1["share_z_available_last"]],
        ["Sensitivität: Median nur berechenbare Quartale", s["median_ccc"], s["median_z"]],
    ]
    put_table(ws, pd.DataFrame(rows, columns=["Messgröße", "CCC", "Z''"]), 4, "Vorlauf")
    for row in (8, 9):
        for col in (2, 3):
            ws.cell(row=row, column=col).number_format = "0%"
    for row in (6, 7, 10):
        for col in (2, 3):
            ws.cell(row=row, column=col).number_format = "0.0"

    # case studies and label audit
    ws = wb.create_sheet("Fallstudien")
    header(ws, "Fallstudien (Auswahl nach fester Regel, docs/case-studies.md)", note)
    cs = pd.read_csv(r / "analysis/case_studies.csv").drop(columns=["company_id", "selection_rule"])
    cs["revenue_ttm_eur"] = cs["revenue_ttm_eur"] / 1e9
    cs = cs.rename(
        columns={
            "case_rank": "Rang",
            "cik": "CIK",
            "company_name": "Unternehmen",
            "sector_de": "Sektor",
            "event_date": "Ereignis",
            "revenue_ttm_eur": "Umsatz TTM Mrd. EUR",
            "lead_months_ccc": "Vorlauf CCC",
            "lead_months_z": "Vorlauf Z''",
        }
    )
    put_table(ws, cs, 4, "Cases", {"Umsatz TTM Mrd. EUR": "0.0", "Vorlauf CCC": "0.0", "Vorlauf Z''": "0.0"})
    la_path = r / "analysis/label_audit.json"
    if la_path.exists():
        la = json.loads(la_path.read_text())
        lab = {
            "filer_is_debtor": "Unternehmen selbst Schuldner",
            "subsidiary_only": "nur Tochter (Wortlaut)",
            "unclear": "nicht erkennbar",
            "item_missing_in_text": "kein Item 1.03 gefunden",
        }
        rows = [[v, la["all_filings"][k], la["first_event_per_company"][k]] for k, v in lab.items()]
        ws.cell(row=10, column=1, value="Prüfung der Labels gegen den Text der 8-K (label_audit.json)").font = Font(
            bold=True
        )
        put_table(ws, pd.DataFrame(rows, columns=["Befund", "alle Filings", "erstes Ereignis"]), 11, "LabelAudit")

    # data quality
    ws = wb.create_sheet("Datenqualität")
    header(ws, "Datenqualitätsregeln (letzter Stand je Regel)", note)
    dq = pd.read_csv(r / "data_quality.csv")
    dq = dq[dq.status != "no_data"].sort_values("run_at").groupby("rule_code").tail(1).sort_values("rule_code")
    dq = dq[
        [
            "rule_code",
            "dq_dimension",
            "table_name",
            "rows_checked",
            "rows_failed",
            "pass_rate",
            "target_pass_rate",
            "status",
        ]
    ]
    put_table(
        ws,
        dq,
        4,
        "DQ",
        {"rows_checked": "#,##0", "rows_failed": "#,##0", "pass_rate": "0.0%", "target_pass_rate": "0.0%"},
    )

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A5"
    for c in wb["Antwort"][4]:
        c.fill, c.font = HEAD, WHITE
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(args.out)


if __name__ == "__main__":
    main()
